"""Tests for the .docx extractor (eval-corpus branch)."""

from pathlib import Path

import pytest
from docx import Document

from backend.app.ingestion.docx_extractor import extract_docx


def _write_docx(path: Path, builder) -> Path:
    doc = Document()
    builder(doc)
    doc.save(str(path))
    return path


def test_single_paragraph(tmp_path: Path):
    def b(doc):
        doc.add_paragraph("The lease term is five years commencing April 2026.")

    p = _write_docx(tmp_path / "lease.docx", b)
    pages = extract_docx(str(p))

    assert len(pages) == 1
    assert pages[0]["page"] == 1
    assert pages[0]["doc_id"] == "lease"
    assert "five years" in pages[0]["text"]


def test_heading_splits_pages(tmp_path: Path):
    def b(doc):
        doc.add_heading("Section 1 — Term", level=1)
        doc.add_paragraph("The agreement is valid for 24 months.")
        doc.add_heading("Section 2 — Payment", level=1)
        doc.add_paragraph("Payable monthly within 7 days of invoice.")

    p = _write_docx(tmp_path / "contract.docx", b)
    pages = extract_docx(str(p))

    assert len(pages) == 2
    assert "24 months" in pages[0]["text"]
    assert "monthly" in pages[1]["text"]


def test_table_extraction(tmp_path: Path):
    def b(doc):
        doc.add_paragraph("Fee schedule below.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Service"
        table.rows[0].cells[1].text = "Amount"
        table.rows[1].cells[0].text = "Setup fee"
        table.rows[1].cells[1].text = "Rs. 15,000"

    p = _write_docx(tmp_path / "quote.docx", b)
    pages = extract_docx(str(p))

    full_text = "\n".join(page["text"] for page in pages)
    assert "Setup fee" in full_text
    assert "15,000" in full_text
    assert "|" in full_text  # table rows are pipe-separated


def test_missing_file_raises():
    with pytest.raises(FileNotFoundError):
        extract_docx("/nonexistent/path.docx")


def test_empty_doc_returns_placeholder(tmp_path: Path):
    def b(doc):
        pass  # empty document

    p = _write_docx(tmp_path / "empty.docx", b)
    pages = extract_docx(str(p))
    assert len(pages) == 1
    assert pages[0]["text"] == ""
